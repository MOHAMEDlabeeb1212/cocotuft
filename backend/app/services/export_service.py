# ==============================================================================
# COCOTUFT PRODUCTION MANAGEMENT SYSTEM - EXCEL EXPORT SERVICE
# ==============================================================================
# Section Purpose: Generates scope-checked, professional .xlsx workbooks using openpyxl.
# Formats editable vs read-only header styling, auto column widths, number formatting,
# frozen headers, and a metadata summary sheet.
# ==============================================================================

import io
from datetime import datetime, date
from typing import List, Dict, Any, Tuple
from sqlalchemy.orm import Session
from sqlalchemy import or_

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.domain import ProductionEntry, ProductionDetail, User, Shift, Machine


ALL_EXPORT_COLUMNS: Dict[str, Tuple[str, bool]] = {
    # key: (Header Title, Is Editable)
    "entry_number": ("Entry No", False),
    "entry_date": ("Entry Date", True),
    "tufted_date": ("Tufted Date", True),
    "shift_name": ("Shift", True),
    "process_name": ("Process", False),
    "machine_name": ("Machine", True),
    "customer_name": ("Customer", True),
    "order_number": ("S.O. Number", True),
    "po_number": ("P.O. Number", True),
    "customer_code": ("Cust. Code", True),
    "roll_number": ("Roll Number", True),
    "base": ("Base Material", True),
    "pile_height": ("Pile Height", True),
    "start_time": ("Start Time", True),
    "end_time": ("End Time", True),
    "length_meters": ("Length (m)", True),
    "width_meters": ("Width (m)", True),
    "actual_qty": ("Actual SQM", False),
    "target_qty": ("Target Qty", True),
    "variation": ("Variation", False),
    "balance_qty": ("Balance Qty", False),
    "belt_speed": ("Belt Speed", True),
    "defects_a_yarn": ("Defect A Yarn", True),
    "defects_b_pvc": ("Defect B PVC", True),
    "defects_c_tufting": ("Defect C Tufting", True),
    "defects_d_stripe": ("Defect D Stripe", True),
    "defects_e_others": ("Defect E Others", True),
    "machine_stop_minutes": ("Machine Stop Mins", True),
    "machine_stop_reason": ("Machine Stop Reason", True),
    "factory_labour_count": ("Factory Labour", True),
    "contract_labour_count": ("Contract Labour", True),
    "status": ("Approval Status", False),
    "worker_name": ("Worker", False),
    "approved_by_name": ("Approved By", False),
    "comments": ("Comments", True),
}


def query_export_records(
    db: Session,
    current_user: User,
    start_date: date = None,
    end_date: date = None,
    shift_ids: List[int] = None,
    worker_ids: List[int] = None,
    status: str = "ALL",
    machine_id: int = None
) -> List[ProductionEntry]:
    """
    Section Purpose: Queries operational records based on user scope and selected filters.
    """
    q = db.query(ProductionEntry)

    # Scope Restrictions: WORKER can only see own; SUPERVISOR/ADMIN see permitted scope
    user_role = current_user.role.role_name.upper() if current_user.role else "WORKER"
    if user_role == "WORKER":
        q = q.filter(ProductionEntry.worker_id == current_user.user_id)
    elif worker_ids and len(worker_ids) > 0:
        q = q.filter(ProductionEntry.worker_id.in_(worker_ids))

    if start_date:
        q = q.filter(ProductionEntry.entry_date >= start_date)
    if end_date:
        q = q.filter(ProductionEntry.entry_date <= end_date)
    if shift_ids and len(shift_ids) > 0:
        q = q.filter(ProductionEntry.shift_id.in_(shift_ids))
    if machine_id:
        q = q.filter(ProductionEntry.machine_id == machine_id)
    if status and status.upper() != "ALL":
        q = q.filter(ProductionEntry.status == status.upper())

    return q.order_by(ProductionEntry.entry_date.desc(), ProductionEntry.entry_id.desc()).all()


def _extract_field_value(entry: ProductionEntry, col_key: str) -> Any:
    detail = entry.details
    if col_key == "entry_number": return entry.entry_number
    if col_key == "entry_date": return entry.entry_date.strftime("%Y-%m-%d") if entry.entry_date else ""
    if col_key == "tufted_date": return entry.tufted_date.strftime("%Y-%m-%d") if entry.tufted_date else ""
    if col_key == "shift_name": return entry.shift.shift_name if entry.shift else ""
    if col_key == "process_name": return entry.process.process_name if entry.process else "TUFTING"
    if col_key == "machine_name": return entry.machine.machine_name if entry.machine else ""
    if col_key == "customer_name": return entry.customer.customer_name if entry.customer else ""
    if col_key == "order_number": return entry.order.order_number if entry.order else ""
    if col_key == "po_number": return detail.po_number if detail else (entry.order.po_number if entry.order else "")
    if col_key == "customer_code": return detail.customer_code if detail else (entry.customer.customer_code if entry.customer else "")
    if col_key == "roll_number": return detail.roll_number if detail else ""
    if col_key == "base": return detail.base if detail else "Natural"
    if col_key == "pile_height": return detail.pile_height if detail else "15 MM"
    if col_key == "start_time": return detail.start_time if detail else ""
    if col_key == "end_time": return detail.end_time if detail else ""
    if col_key == "length_meters": return detail.length_meters if detail else 0.0
    if col_key == "width_meters": return detail.width_meters if detail else 0.0
    if col_key == "actual_qty": return detail.actual_qty if detail else 0.0
    if col_key == "target_qty": return detail.target_qty if detail else 500.0
    if col_key == "variation": return detail.variation if detail else 0.0
    if col_key == "balance_qty": return detail.balance_qty if detail else 0.0
    if col_key == "belt_speed": return detail.belt_speed if detail else "Normal"
    if col_key == "defects_a_yarn": return detail.defects_a_yarn if detail else 0
    if col_key == "defects_b_pvc": return detail.defects_b_pvc if detail else 0
    if col_key == "defects_c_tufting": return detail.defects_c_tufting if detail else 0
    if col_key == "defects_d_stripe": return detail.defects_d_stripe if detail else 0
    if col_key == "defects_e_others": return detail.defects_e_others if detail else 0
    if col_key == "machine_stop_minutes": return detail.machine_stop_minutes if detail else 0
    if col_key == "machine_stop_reason": return detail.machine_stop_reason if detail else ""
    if col_key == "factory_labour_count": return detail.factory_labour_count if detail else 1
    if col_key == "contract_labour_count": return detail.contract_labour_count if detail else 0
    if col_key == "status": return entry.status.value if hasattr(entry.status, 'value') else str(entry.status)
    if col_key == "worker_name": return entry.worker.full_name if entry.worker else ""
    if col_key == "approved_by_name": return entry.approved_by.full_name if entry.approved_by else ""
    if col_key == "comments": return detail.comments if detail else ""
    return ""


def generate_excel_workbook(
    db: Session,
    current_user: User,
    entries: List[ProductionEntry],
    selected_columns: List[str] = None,
    filter_summary: Dict[str, Any] = None
) -> bytes:
    """
    Section Purpose: Builds openpyxl workbook binary stream with Production Data and Export Summary sheets.
    """
    wb = openpyxl.Workbook()
    
    # --------------------------------------------------------------------------
    # Sheet 1: Production Records
    # --------------------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Production Records"

    cols_to_export = [c for c in (selected_columns or list(ALL_EXPORT_COLUMNS.keys())) if c in ALL_EXPORT_COLUMNS]
    if not cols_to_export:
        cols_to_export = list(ALL_EXPORT_COLUMNS.keys())

    # Styling definitions
    font_bold = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Arial", size=10)
    
    fill_editable_header = PatternFill(start_color="203764", end_color="203764", fill_type="solid") # Dark Navy
    fill_readonly_header = PatternFill(start_color="595959", end_color="595959", fill_type="solid") # Dark Gray
    
    fill_editable_cell = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid") # Soft Blue Tint
    fill_readonly_cell = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # Soft Gray

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Header Row
    for col_idx, col_key in enumerate(cols_to_export, start=1):
        title, is_editable = ALL_EXPORT_COLUMNS[col_key]
        display_title = f"{title} *" if is_editable else f"{title} (RO)"
        cell = ws1.cell(row=1, column=col_idx, value=display_title)
        cell.font = font_bold
        cell.fill = fill_editable_header if is_editable else fill_readonly_header
        cell.alignment = align_center
        cell.border = thin_border

    ws1.freeze_panes = "A2"
    ws1.row_dimensions[1].height = 28

    # Data Rows
    for row_idx, entry in enumerate(entries, start=2):
        ws1.row_dimensions[row_idx].height = 20
        for col_idx, col_key in enumerate(cols_to_export, start=1):
            val = _extract_field_value(entry, col_key)
            _, is_editable = ALL_EXPORT_COLUMNS[col_key]
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_data
            cell.border = thin_border
            cell.fill = fill_editable_cell if is_editable else fill_readonly_cell

            # Number formatting & Alignment
            if isinstance(val, float):
                cell.number_format = '0.00'
                cell.alignment = align_right
            elif isinstance(val, int):
                cell.number_format = '#,##0'
                cell.alignment = align_center
            elif col_key in ["entry_date", "tufted_date", "start_time", "end_time", "status"]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Auto column width adjustment
    for col_idx, col_key in enumerate(cols_to_export, start=1):
        max_len = len(ALL_EXPORT_COLUMNS[col_key][0]) + 5
        col_letter = get_column_letter(col_idx)
        for row in range(2, len(entries) + 2):
            cell_val = str(ws1.cell(row=row, column=col_idx).value or "")
            if len(cell_val) > max_len:
                max_len = len(cell_val)
        ws1.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    # --------------------------------------------------------------------------
    # Sheet 2: Export Summary Metadata
    # --------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="Export Summary")
    ws2.column_dimensions["A"].width = 25
    ws2.column_dimensions["B"].width = 50

    ws2.cell(row=1, column=1, value="COCOTUFT PRODUCTION SYSTEM - EXPORT METADATA").font = Font(name="Arial", size=12, bold=True, color="203764")
    
    summary_items = [
        ("Export Timestamp", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")),
        ("Exported By User", f"{current_user.full_name} ({current_user.username})"),
        ("User Role", current_user.role.role_name if current_user.role else "WORKER"),
        ("Total Records Exported", len(entries)),
        ("Date Range Filter", f"{filter_summary.get('start_date', 'All')} to {filter_summary.get('end_date', 'All')}" if filter_summary else "All"),
        ("Shifts Filter", filter_summary.get('shifts', 'All') if filter_summary else "All"),
        ("Status Filter", filter_summary.get('status', 'ALL') if filter_summary else "ALL"),
        ("Legend", "* Header indicates Editable Field | (RO) Header indicates Read-Only Field"),
    ]

    for idx, (label, val) in enumerate(summary_items, start=3):
        c1 = ws2.cell(row=idx, column=1, value=label)
        c1.font = Font(name="Arial", size=10, bold=True)
        c2 = ws2.cell(row=idx, column=2, value=val)
        c2.font = Font(name="Arial", size=10)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
